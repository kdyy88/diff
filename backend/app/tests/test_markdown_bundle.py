from pathlib import Path
import tempfile

from docx import Document as DocxDocument

from app.models.schemas import ChapterDiffSummary, DiffAnchor, DiffResult, DiffSummary, HighlightFragment, PageMeta, TableContext
from app.services.differ import compare_documents
from app.services.extractor import extract_document
from app.services.markdown_bundle import MISSING_SIDE_PLACEHOLDER, build_markdown_bundle


def _fragment(page: int) -> HighlightFragment:
    return HighlightFragment(page=page, bbox=[10.0, 10.0, 20.0, 20.0], viewport_ref=f"page-{page}")


def _build_result(
    anchors: list[DiffAnchor],
    *,
    chapters: list[ChapterDiffSummary] | None = None,
    document_kind: str = "pdf",
) -> DiffResult:
    return DiffResult(
        document_kind=document_kind,
        pages_left=[PageMeta(page=index, width=600.0, height=800.0) for index in range(3)],
        pages_right=[PageMeta(page=index, width=600.0, height=800.0) for index in range(3)],
        summary=DiffSummary(
            insertions=sum(1 for anchor in anchors if anchor.kind == "insert"),
            deletions=sum(1 for anchor in anchors if anchor.kind == "delete"),
            replacements=sum(1 for anchor in anchors if anchor.kind == "replace"),
            reflows=sum(1 for anchor in anchors if anchor.kind == "reflow"),
            pages_a=3,
            pages_b=3,
        ),
        anchors=anchors,
        chapters=chapters or [],
    )


def _compare_docx(left_path: Path, right_path: Path) -> DiffResult:
    left = extract_document(left_path, header_margin=0, footer_margin=0)
    right = extract_document(right_path, header_margin=0, footer_margin=0)
    return compare_documents(left, right, include_reflow=False)


def test_build_markdown_bundle_uses_part_mode_and_missing_side_placeholder() -> None:
    result = _build_result(
        [
            DiffAnchor(
                id="anchor-insert-1",
                kind="insert",
                excerpt_left="",
                excerpt_right="新增了一段带上下文的内容。",
                right_fragments=[_fragment(1)],
            ),
            DiffAnchor(
                id="reflow-1",
                kind="reflow",
                excerpt_left="layout moved",
                excerpt_right="layout moved",
                left_fragments=[_fragment(0)],
                right_fragments=[_fragment(2)],
            ),
        ]
    )

    bundle = build_markdown_bundle("job-123", result)
    detail = bundle.files[1].content

    assert [item.name for item in bundle.files] == ["00-overview.md", "10-part-01.md"]
    assert MISSING_SIDE_PLACEHOLDER not in detail
    assert "（原文档无对应内容）" in detail
    assert "原文档页码：无" in detail
    assert "修订后文档页码：第 2 页" in detail
    assert "reflow-1" not in detail
    assert "reflow：1" in bundle.files[0].content


def test_build_markdown_bundle_falls_back_when_chapter_mapping_is_incomplete() -> None:
    result = _build_result(
        [
            DiffAnchor(
                id="anchor-1",
                kind="replace",
                excerpt_left="旧文本",
                excerpt_right="新文本",
                left_fragments=[_fragment(0)],
                right_fragments=[_fragment(0)],
                chapter_id="missing-chapter",
            )
        ],
        chapters=[
            ChapterDiffSummary(
                id="known-chapter",
                title="Chapter 1",
                index=0,
                anchor_count=1,
                first_anchor_id="anchor-1",
                summary=DiffSummary(replacements=1, pages_a=1, pages_b=1),
            )
        ],
    )

    bundle = build_markdown_bundle("job-123", result)

    assert [item.name for item in bundle.files] == ["00-overview.md", "10-part-01.md"]
    assert "## 分片索引" in bundle.files[0].content


def test_build_markdown_bundle_splits_by_chapter_and_is_stable() -> None:
    anchors = [
        DiffAnchor(
            id=f"chapter-a-anchor-{index:02d}",
            kind="replace",
            excerpt_left=f"旧内容 {index}",
            excerpt_right=f"新内容 {index}",
            left_fragments=[_fragment(0)],
            right_fragments=[_fragment(0)],
            chapter_id="chapter-a",
            chapter_title="Chapter A",
            chapter_index=0,
        )
        for index in range(21)
    ]
    result = _build_result(
        anchors,
        chapters=[
            ChapterDiffSummary(
                id="chapter-a",
                title="Chapter A",
                index=0,
                anchor_count=21,
                first_anchor_id="chapter-a-anchor-00",
                summary=DiffSummary(replacements=21, pages_a=1, pages_b=1),
            ),
            ChapterDiffSummary(
                id="chapter-b",
                title="Chapter B",
                index=1,
                anchor_count=0,
                first_anchor_id=None,
                summary=DiffSummary(pages_a=1, pages_b=1),
            ),
        ],
    )

    bundle_one = build_markdown_bundle("job-123", result)
    bundle_two = build_markdown_bundle("job-123", result)

    assert [item.name for item in bundle_one.files] == [
        "00-overview.md",
        "10-chapter-01-part-01.md",
        "10-chapter-01-part-02.md",
    ]
    assert [(item.name, item.content) for item in bundle_one.files] == [
        (item.name, item.content) for item in bundle_two.files
    ]
    assert "Chapter B（anchor: 0，详情：无 detail shard）" in bundle_one.files[0].content
    assert "anchor_count: 20" in bundle_one.files[1].content
    assert "anchor_count: 1" in bundle_one.files[2].content


def test_build_markdown_bundle_renders_table_changes_for_llm_consumption() -> None:
    result = _build_result(
        [
            DiffAnchor(
                id="table-cell-1",
                kind="replace",
                source_type="table",
                excerpt_left="1.2",
                excerpt_right="0.8",
                left_fragments=[_fragment(0)],
                right_fragments=[_fragment(0)],
                raw_event_count=2,
                confidence="low",
                table_context=TableContext(
                    table_id="table-1",
                    row=1,
                    col=0,
                    row_label="化合物 Zanubrutinib",
                    col_label="靶点亲和力",
                ),
            ),
            DiffAnchor(
                id="table-cell-2",
                kind="replace",
                source_type="table",
                excerpt_left="旧值",
                excerpt_right="新值",
                left_fragments=[_fragment(1)],
                right_fragments=[_fragment(1)],
                table_context=TableContext(table_id="table-1", row=2, col=1),
            ),
            DiffAnchor(
                id="table-structure-1",
                kind="replace",
                source_type="table",
                excerpt_left="Table structure changed | 3x3",
                excerpt_right="Table structure changed | 4x3",
                left_fragments=[_fragment(2)],
                right_fragments=[_fragment(2)],
                table_context=TableContext(table_id="table-1", row=-1, col=-1),
            ),
        ]
    )

    bundle = build_markdown_bundle("job-123", result)
    overview = bundle.files[0].content
    detail = bundle.files[1].content

    assert "表格变更：所在列 [靶点亲和力]，所在行 [化合物 Zanubrutinib]，原值为 \"1.2\"，变更为 \"0.8\"。" in detail
    assert "表格变更：所在列 [C2]，所在行 [R3]，原值为 \"旧值\"，变更为 \"新值\"。" in detail
    assert "表格结构发生变化。" in detail
    assert "原文档页码：第 1 页" in detail
    assert "修订后文档页码：第 1 页" in detail
    assert "low confidence" in overview
    assert "contains_tables: true" in detail


def test_build_markdown_bundle_uses_explicit_placeholder_for_table_delete() -> None:
    result = _build_result(
        [
            DiffAnchor(
                id="table-delete-1",
                kind="delete",
                source_type="table",
                excerpt_left="1.2",
                excerpt_right="",
                left_fragments=[_fragment(0)],
                table_context=TableContext(
                    table_id="table-1",
                    row=0,
                    col=0,
                    row_label="化合物 Zanubrutinib",
                    col_label="靶点亲和力",
                ),
            )
        ]
    )

    bundle = build_markdown_bundle("job-123", result)
    detail = bundle.files[1].content

    assert "表格变更：所在列 [靶点亲和力]，所在行 [化合物 Zanubrutinib]，原值为 \"1.2\"，变更为 \"（修订后文档无对应内容）\"。" in detail


def test_build_markdown_bundle_uses_document_role_labels_in_overview_and_detail() -> None:
    result = _build_result(
        [
            DiffAnchor(
                id="anchor-1",
                kind="replace",
                excerpt_left="旧文本",
                excerpt_right="新文本",
                left_fragments=[_fragment(0)],
                right_fragments=[_fragment(1)],
            )
        ]
    )

    bundle = build_markdown_bundle("job-123", result)
    overview = bundle.files[0].content
    detail = bundle.files[1].content

    assert 'original_document_label: "原文档"' in overview
    assert 'modified_document_label: "修订后文档"' in detail
    assert "原文档：旧文本 | 修订后文档：新文本" in overview
    assert "原文档片段：旧文本" in detail
    assert "修订后文档片段：新文本" in detail


def test_build_markdown_bundle_renders_docx_text_diff_details() -> None:
    left_path = Path(tempfile.gettempdir()) / "bundle-left.docx"
    right_path = Path(tempfile.gettempdir()) / "bundle-right.docx"

    left_doc = DocxDocument()
    left_doc.add_heading("概述", level=1)
    left_doc.add_paragraph("该药物对肿瘤细胞增殖具有明显抑制作用，且安全性良好。")
    left_doc.save(left_path)

    right_doc = DocxDocument()
    right_doc.add_heading("概述", level=1)
    right_doc.add_paragraph("该药物对肿瘤细胞增殖具有明显促进作用，且安全性良好。")
    right_doc.save(right_path)

    result = _compare_docx(left_path, right_path)
    bundle = build_markdown_bundle(
        "job-docx-1",
        result,
        source_document_path=left_path,
        modified_document_path=right_path,
    )
    detail = bundle.files[1].content

    assert "块类型：段落" in detail
    assert "原文档位置：节点 paragraph-1" in detail
    assert "修订后文档位置：节点 paragraph-1" in detail
    assert "原文档全文：该药物对肿瘤细胞增殖具有明显抑制作用，且安全性良好。" in detail
    assert "修订后文档全文：该药物对肿瘤细胞增殖具有明显促进作用，且安全性良好。" in detail
    assert "差异明细：" in detail
    assert "替换：" in detail
    assert "抑制" in detail
    assert "促进" in detail


def test_build_markdown_bundle_renders_docx_insert_with_placeholder() -> None:
    left_path = Path(tempfile.gettempdir()) / "bundle-insert-left.docx"
    right_path = Path(tempfile.gettempdir()) / "bundle-insert-right.docx"

    left_doc = DocxDocument()
    left_doc.add_heading("概述", level=1)
    left_doc.add_paragraph("现有内容。")
    left_doc.save(left_path)

    right_doc = DocxDocument()
    right_doc.add_heading("概述", level=1)
    right_doc.add_paragraph("现有内容。")
    right_doc.add_paragraph("新增的风险说明。")
    right_doc.save(right_path)

    result = _compare_docx(left_path, right_path)
    bundle = build_markdown_bundle(
        "job-docx-2",
        result,
        source_document_path=left_path,
        modified_document_path=right_path,
    )
    detail = bundle.files[1].content

    assert "原文档全文：（原文档无对应内容）" in detail
    assert "修订后文档全文：新增的风险说明。" in detail
    assert "- 新增：\"新增的风险说明。\"" in detail


def test_build_markdown_bundle_marks_docx_moved_paragraph_as_position_change() -> None:
    left_path = Path(tempfile.gettempdir()) / "bundle-move-left.docx"
    right_path = Path(tempfile.gettempdir()) / "bundle-move-right.docx"

    moved_paragraph = "- 提取层拆成适配器：PDF 继续走现有 extractor/projector；DOCX 新增 DocxProjection。"

    left_doc = DocxDocument()
    left_doc.add_paragraph("摘要")
    left_doc.add_paragraph("保留现有 PDF 链路。")
    left_doc.add_paragraph(moved_paragraph)
    left_doc.add_paragraph("接口保持兼容。")
    left_doc.add_paragraph("测试覆盖需要补齐。")
    left_doc.save(left_path)

    right_doc = DocxDocument()
    right_doc.add_paragraph("摘要")
    right_doc.add_paragraph("保留现有 PDF 链路。")
    right_doc.add_paragraph("接口保持兼容。")
    right_doc.add_paragraph("测试覆盖需要补齐。")
    right_doc.add_paragraph(moved_paragraph)
    right_doc.save(right_path)

    result = _compare_docx(left_path, right_path)
    bundle = build_markdown_bundle(
        "job-docx-move",
        result,
        source_document_path=left_path,
        modified_document_path=right_path,
    )
    detail = bundle.files[1].content

    assert "位置移动：文本内容未改写" in detail
    assert "与 anchor-text-1 配对后判定为移动" in detail or "与 anchor-text-0 配对后判定为移动" in detail
    assert "原文档位置：节点 paragraph-2[0:" in detail
    assert "修订后文档位置：节点 paragraph-4[0:" in detail
    assert "原文档全文：- 提取层拆成适配器：PDF 继续走现有 extractor/projector；DOCX 新增 DocxProjection。" in detail
    assert "修订后文档全文：- 提取层拆成适配器：PDF 继续走现有 extractor/projector；DOCX 新增 DocxProjection。" in detail


def test_build_markdown_bundle_renders_docx_table_diff_details() -> None:
    left_path = Path(tempfile.gettempdir()) / "bundle-table-left.docx"
    right_path = Path(tempfile.gettempdir()) / "bundle-table-right.docx"

    left_doc = DocxDocument()
    left_table = left_doc.add_table(rows=2, cols=2)
    left_table.cell(0, 0).text = "药物"
    left_table.cell(0, 1).text = "结果"
    left_table.cell(1, 0).text = "Zanubrutinib"
    left_table.cell(1, 1).text = "1.2"
    left_doc.save(left_path)

    right_doc = DocxDocument()
    right_table = right_doc.add_table(rows=2, cols=2)
    right_table.cell(0, 0).text = "药物"
    right_table.cell(0, 1).text = "结果"
    right_table.cell(1, 0).text = "Zanubrutinib"
    right_table.cell(1, 1).text = "0.8"
    right_doc.save(right_path)

    result = _compare_docx(left_path, right_path)
    bundle = build_markdown_bundle(
        "job-docx-3",
        result,
        source_document_path=left_path,
        modified_document_path=right_path,
    )
    detail = bundle.files[1].content

    assert "块类型：表格单元格" in detail
    assert "表格变更：所在列 [结果]，所在行 [Zanubrutinib]，原值为 \"1.2\"，变更为 \"0.8\"。" in detail
    assert "原文档全文：1.2" in detail
    assert "修订后文档全文：0.8" in detail
    assert "替换：\"1.2\" -> \"0.8\"" in detail