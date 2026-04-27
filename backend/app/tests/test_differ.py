from pathlib import Path
import tempfile

from docx import Document as DocxDocument
import fitz

from app.services.differ import compare_documents
from app.services.extractor import CharAtom, DocumentProjection, PageInfo, TextSegment, WordAtom, extract_document


def _collapsed_mapping(text: str) -> tuple[str, list[int]]:
    aligned_chars = []
    aligned_to_raw = []
    previous_was_space = False
    for index, char in enumerate(text):
        normalized = " " if char.isspace() else char
        if normalized.isspace():
            if previous_was_space:
                continue
            aligned_chars.append(" ")
            aligned_to_raw.append(index)
            previous_was_space = True
            continue
        aligned_chars.append(normalized)
        aligned_to_raw.append(index)
        previous_was_space = False
    return "".join(aligned_chars), aligned_to_raw


def _build_document(text: str, *, page: int = 0, y: float = 100.0) -> DocumentProjection:
    chars = []
    words = []
    cursor = 0
    current_x = 20.0
    for word_index, token in enumerate(text.split(" ")):
        char_start = len(chars)
        for char in token:
            box = (current_x, y, current_x + 8.0, y + 12.0)
            chars.append(
                CharAtom(
                    id=f"c-{len(chars)}",
                    char=char,
                    norm_char=char,
                    page=page,
                    bbox=box,
                    block=0,
                    line=0,
                    word=word_index,
                    stream_index=len(chars),
                    synthetic=False,
                )
            )
            current_x += 8.0
        char_end = len(chars)
        words.append(
            WordAtom(
                id=f"w-{word_index}",
                text=token,
                norm_text=token,
                page=page,
                bbox=(current_x - (len(token) * 8.0), y, current_x, y + 12.0),
                block=0,
                line=0,
                word=word_index,
                char_start=char_start,
                char_end=char_end,
            )
        )
        if word_index < len(text.split(" ")) - 1:
            chars.append(
                CharAtom(
                    id=f"c-{len(chars)}",
                    char=" ",
                    norm_char=" ",
                    page=page,
                    bbox=None,
                    block=0,
                    line=0,
                    word=word_index,
                    stream_index=len(chars),
                    synthetic=True,
                )
            )
            current_x += 4.0
    return DocumentProjection(
        pdf_path=None,  # type: ignore[arg-type]
        document_kind="pdf",
        pages=[PageInfo(page=page, width=600.0, height=800.0)],
        chars=chars,
        words=words,
        tables=[],
        text_segments=[
            TextSegment(
                id="segment-0",
                page=page,
                block=0,
                line=0,
                bbox=(20.0, y, current_x, y + 12.0),
                raw_text=text,
                aligned_text=" ".join(text.split()),
                raw_start=0,
                raw_end=len(chars),
            )
        ],
        raw_text=text,
        normalized_text=text,
        aligned_text=" ".join(text.split()),
        aligned_to_raw=[index for index, char in enumerate(text) if (not char.isspace()) or (index == 0 or not text[index - 1].isspace())],
        review_html=None,
    )


def _build_character_document(text: str, *, page: int = 0, y: float = 100.0) -> DocumentProjection:
    chars = []
    current_x = 20.0
    for index, char in enumerate(text):
        bbox = None if char.isspace() else (current_x, y, current_x + 8.0, y + 12.0)
        chars.append(
            CharAtom(
                id=f"c-{index}",
                char=char,
                norm_char=" " if char.isspace() else char,
                page=page,
                bbox=bbox,
                block=0,
                line=0,
                word=None,
                stream_index=index,
                synthetic=char.isspace(),
            )
        )
        current_x += 8.0
    return DocumentProjection(
        pdf_path=None,  # type: ignore[arg-type]
        document_kind="pdf",
        pages=[PageInfo(page=page, width=600.0, height=800.0)],
        chars=chars,
        words=[],
        tables=[],
        text_segments=[
            TextSegment(
                id="segment-0",
                page=page,
                block=0,
                line=0,
                bbox=(20.0, y, current_x, y + 12.0),
                raw_text=text,
                aligned_text=" ".join(text.split()),
                raw_start=0,
                raw_end=len(chars),
            )
        ],
        raw_text=text,
        normalized_text="".join(" " if char.isspace() else char for char in text),
        aligned_text=" ".join(text.split()),
        aligned_to_raw=[index for index, char in enumerate(text) if (not char.isspace()) or (index == 0 or not text[index - 1].isspace())],
        review_html=None,
    )


def _build_wrapped_document(lines: list[str], *, page: int = 0, y: float = 100.0) -> DocumentProjection:
    chars = []
    current_y = y
    raw_parts: list[str] = []
    text_segments: list[TextSegment] = []
    for line_index, line in enumerate(lines):
        current_x = 20.0
        segment_start = len(chars)
        for char in line:
            stream_index = len(chars)
            chars.append(
                CharAtom(
                    id=f"c-{stream_index}",
                    char=char,
                    norm_char=" " if char.isspace() else char,
                    page=page,
                    bbox=None if char.isspace() else (current_x, current_y, current_x + 8.0, current_y + 12.0),
                    block=0,
                    line=line_index,
                    word=None,
                    stream_index=stream_index,
                    synthetic=False,
                )
            )
            current_x += 8.0
            raw_parts.append(char)

        segment_end = len(chars)
        segment_aligned_text, _ = _collapsed_mapping(line)
        text_segments.append(
            TextSegment(
                id=f"segment-{line_index}",
                page=page,
                block=0,
                line=line_index,
                bbox=(20.0, current_y, current_x, current_y + 12.0),
                raw_text=line,
                aligned_text=segment_aligned_text,
                raw_start=segment_start,
                raw_end=segment_end,
            )
        )

        if line_index < len(lines) - 1:
            stream_index = len(chars)
            chars.append(
                CharAtom(
                    id=f"c-{stream_index}",
                    char="\n",
                    norm_char=" ",
                    page=page,
                    bbox=None,
                    block=0,
                    line=line_index,
                    word=None,
                    stream_index=stream_index,
                    synthetic=True,
                )
            )
            raw_parts.append("\n")
            current_y += 16.0

    aligned_chars = []
    aligned_to_raw = []
    previous_was_space = False
    for raw_index, char in enumerate(chars):
        if char.synthetic and char.char.isspace():
            continue
        normalized = char.norm_char or char.char
        if normalized.isspace():
            if previous_was_space:
                continue
            aligned_chars.append(" ")
            aligned_to_raw.append(raw_index)
            previous_was_space = True
        else:
            aligned_chars.append(normalized)
            aligned_to_raw.append(raw_index)
            previous_was_space = False

    return DocumentProjection(
        pdf_path=None,  # type: ignore[arg-type]
        document_kind="pdf",
        pages=[PageInfo(page=page, width=600.0, height=800.0)],
        chars=chars,
        words=[],
        tables=[],
        text_segments=text_segments,
        raw_text="".join(raw_parts),
        normalized_text="".join(char.norm_char for char in chars),
        aligned_text="".join(aligned_chars),
        aligned_to_raw=aligned_to_raw,
        review_html=None,
    )


def _build_real_space_document(text: str, *, page: int = 0, y: float = 100.0) -> DocumentProjection:
    chars = []
    current_x = 20.0
    for index, char in enumerate(text):
        bbox = None if char.isspace() else (current_x, y, current_x + 8.0, y + 12.0)
        chars.append(
            CharAtom(
                id=f"c-{index}",
                char=char,
                norm_char=" " if char.isspace() else char,
                page=page,
                bbox=bbox,
                block=0,
                line=0,
                word=None,
                stream_index=index,
                synthetic=False,
            )
        )
        current_x += 8.0

    aligned_chars = []
    aligned_to_raw = []
    previous_was_space = False
    for raw_index, char in enumerate(chars):
        normalized = char.norm_char or char.char
        if normalized.isspace():
            if previous_was_space:
                continue
            aligned_chars.append(" ")
            aligned_to_raw.append(raw_index)
            previous_was_space = True
        else:
            aligned_chars.append(normalized)
            aligned_to_raw.append(raw_index)
            previous_was_space = False

    return DocumentProjection(
        pdf_path=None,  # type: ignore[arg-type]
        document_kind="pdf",
        pages=[PageInfo(page=page, width=600.0, height=800.0)],
        chars=chars,
        words=[],
        tables=[],
        text_segments=[
            TextSegment(
                id="segment-0",
                page=page,
                block=0,
                line=0,
                bbox=(20.0, y, current_x, y + 12.0),
                raw_text=text,
                aligned_text="".join(aligned_chars),
                raw_start=0,
                raw_end=len(chars),
            )
        ],
        raw_text=text,
        normalized_text="".join(char.norm_char for char in chars),
        aligned_text="".join(aligned_chars),
        aligned_to_raw=aligned_to_raw,
        review_html=None,
    )


def test_compare_documents_emits_word_dom_fragments_for_docx_replace() -> None:
    left_path = Path(tempfile.gettempdir()) / "differ-left.docx"
    right_path = Path(tempfile.gettempdir()) / "differ-right.docx"

    left_doc = DocxDocument()
    left_doc.add_paragraph("Chapter 1", style="Heading 1")
    left_doc.add_paragraph("Original body text")
    left_table = left_doc.add_table(rows=1, cols=2)
    left_table.cell(0, 0).text = "A1"
    left_table.cell(0, 1).text = "B1"
    left_doc.save(left_path)

    right_doc = DocxDocument()
    right_doc.add_paragraph("Chapter 1", style="Heading 1")
    right_doc.add_paragraph("Updated body text")
    right_table = right_doc.add_table(rows=1, cols=2)
    right_table.cell(0, 0).text = "A1"
    right_table.cell(0, 1).text = "B2"
    right_doc.save(right_path)

    left = extract_document(left_path, header_margin=0, footer_margin=0)
    right = extract_document(right_path, header_margin=0, footer_margin=0)
    result = compare_documents(left, right, include_reflow=True)

    assert result.document_kind == "docx"
    assert result.summary.reflows == 0
    text_anchor = next(anchor for anchor in result.anchors if anchor.source_type == "text")
    assert text_anchor.left_fragments[0].kind == "word"
    assert text_anchor.left_fragments[0].dom_id == "paragraph-1"
    table_anchor = next(anchor for anchor in result.anchors if anchor.source_type == "table")
    assert table_anchor.right_fragments[0].kind == "word"
    assert table_anchor.right_fragments[0].dom_id == "table-2-r0-c1"


def test_compare_documents_orders_docx_anchors_by_dom_position() -> None:
    left_path = Path(tempfile.gettempdir()) / "differ-order-left.docx"
    right_path = Path(tempfile.gettempdir()) / "differ-order-right.docx"

    left_doc = DocxDocument()
    left_doc.add_paragraph("Chapter 1", style="Heading 1")
    left_doc.add_paragraph("Alpha prefix with a long stable lead ending in cat")
    left_doc.add_paragraph("Beta dog")
    left_doc.add_paragraph("Gamma yak")
    left_doc.save(left_path)

    right_doc = DocxDocument()
    right_doc.add_paragraph("Chapter 1", style="Heading 1")
    right_doc.add_paragraph("Alpha prefix with a long stable lead ending in dog")
    right_doc.add_paragraph("Beta fog")
    right_doc.add_paragraph("Gamma zap")
    right_doc.save(right_path)

    left = extract_document(left_path, header_margin=0, footer_margin=0)
    right = extract_document(right_path, header_margin=0, footer_margin=0)
    result = compare_documents(left, right, include_reflow=False)

    text_anchor_dom_ids = [
        anchor.left_fragments[0].dom_id
        for anchor in result.anchors
        if anchor.source_type == "text" and anchor.left_fragments and anchor.left_fragments[0].kind == "word"
    ]

    assert text_anchor_dom_ids == ["paragraph-1", "paragraph-2", "paragraph-3"]


def _build_table_pdf(path: Path, cell_values: list[list[str]], *, top_label: str) -> None:
    doc = fitz.open()
    row_height = 36
    table_top = 60
    table_bottom = table_top + (row_height * len(cell_values))
    page = doc.new_page(width=400, height=max(260, table_bottom + 80))
    page.insert_text((24, 24), top_label)

    x_positions = [24, 124, 244, 364]
    y_positions = [table_top + (row_height * index) for index in range(len(cell_values) + 1)]
    for x in x_positions:
        page.draw_line((x, y_positions[0]), (x, y_positions[-1]), color=(0, 0, 0), width=1)
    for y in y_positions:
        page.draw_line((x_positions[0], y), (x_positions[-1], y), color=(0, 0, 0), width=1)

    for row_index, row in enumerate(cell_values):
        for col_index, value in enumerate(row):
            page.insert_text((x_positions[col_index] + 8, y_positions[row_index] + 22), value)

    doc.save(path)
    doc.close()


def test_compare_documents_emits_replace_anchor() -> None:
    left = _build_document("Dose 0.5 mg")
    right = _build_document("Dose 0.05 mg")

    result = compare_documents(left, right, include_reflow=False)

    assert result.summary.replacements == 1
    assert any(anchor.kind == "replace" for anchor in result.anchors)


def test_compare_documents_uses_contextual_excerpts_for_text_replace() -> None:
    left = _build_document("The subject received 0.5 mg daily; adverse events remained stable")
    right = _build_document("The subject received 0.05 mg daily; adverse events remained stable")

    result = compare_documents(left, right, include_reflow=False)
    anchor = next(anchor for anchor in result.anchors if anchor.kind == "replace")

    assert "The subject received" in anchor.excerpt_left
    assert "mg daily" in anchor.excerpt_left
    assert "The subject received" in anchor.excerpt_right
    assert len(anchor.excerpt_left) <= 160
    assert len(anchor.excerpt_right) <= 160


def test_compare_documents_avoids_isolated_word_excerpts_for_replace() -> None:
    left = _build_document("该药物对肿瘤细胞增殖具有明显抑制作用，且安全性良好。")
    right = _build_document("该药物对肿瘤细胞增殖具有明显促进作用，且安全性良好。")

    result = compare_documents(left, right, include_reflow=False)
    anchor = next(anchor for anchor in result.anchors if anchor.kind == "replace")

    assert anchor.excerpt_left != "抑制"
    assert anchor.excerpt_right != "促进"
    assert "肿瘤细胞增殖" in anchor.excerpt_left
    assert "肿瘤细胞增殖" in anchor.excerpt_right
    assert "作用" in anchor.excerpt_left
    assert "作用" in anchor.excerpt_right


def test_compare_documents_emits_reflow_anchor_for_large_equal_block() -> None:
    text = "This section remains stable across pages but moves later into the next page without content changes"
    left = _build_document(text, page=0, y=120.0)
    right = _build_document(text, page=1, y=60.0)

    result = compare_documents(left, right, include_reflow=True)

    assert result.summary.reflows == 1
    assert result.anchors[0].kind == "reflow"


def test_compare_documents_coalesces_overlapping_replace_anchors() -> None:
    left = _build_document("超时未确认的，视为承租人认可")
    right = _build_document("确认、超时未确认，视为承租人认可")

    result = compare_documents(left, right, include_reflow=False)
    replace_anchors = [anchor for anchor in result.anchors if anchor.kind == "replace"]

    assert len(replace_anchors) == 1
    assert replace_anchors[0].raw_event_count >= 1
    assert "确认" in replace_anchors[0].excerpt_right


def test_compare_documents_does_not_merge_across_strong_boundaries() -> None:
    left = _build_document("第一句改动。第二句保持")
    right = _build_document("第一句变化。第二句调整")

    result = compare_documents(left, right, include_reflow=False)
    replace_anchors = [anchor for anchor in result.anchors if anchor.kind == "replace"]

    assert len(replace_anchors) >= 2


def test_compare_documents_ignores_whitespace_only_changes() -> None:
    left = _build_character_document("因出租人或自如原因影响承租人居住安全。")
    right = _build_character_document("因出租人或自如原因影响承租人 居住安全。")

    result = compare_documents(left, right, include_reflow=False)

    assert result.summary.insertions == 0
    assert result.summary.deletions == 0
    assert result.summary.replacements == 0
    assert result.anchors == []


def test_compare_documents_ignores_synthetic_line_wrap_differences() -> None:
    left = _build_wrapped_document(["如本合同为续约合同，支付时间不受本条款限制，以合同内的具体约", "定为准"])
    right = _build_wrapped_document(["如本合同为续约合同，支付时间不受本条款限制，以合同内的具体约定为", "准"])

    result = compare_documents(left, right, include_reflow=False)

    assert result.summary.insertions == 0
    assert result.summary.deletions == 0
    assert result.summary.replacements == 0
    assert result.anchors == []


def test_compare_documents_preserves_real_space_differences() -> None:
    left = _build_real_space_document("A B")
    right = _build_real_space_document("AB")

    result = compare_documents(left, right, include_reflow=False)

    assert result.summary.insertions + result.summary.deletions + result.summary.replacements > 0
    assert result.anchors


def test_compare_documents_reanchors_after_inserted_lines() -> None:
    left = _build_wrapped_document(["A", "B", "C", "Tail stable"])
    right = _build_wrapped_document(["A", "Inserted one", "Inserted two", "B", "C", "Tail stable"])

    result = compare_documents(left, right, include_reflow=False)

    assert all("Tail stable" not in anchor.excerpt_left for anchor in result.anchors)
    assert all("Tail stable" not in anchor.excerpt_right for anchor in result.anchors)


def test_compare_documents_marks_large_replace_windows_low_confidence() -> None:
    left = _build_wrapped_document(["Prefix", "L1", "L2", "L3", "L4", "L5", "Suffix"])
    right = _build_wrapped_document(["Prefix", "R1", "R2", "R3", "R4", "R5", "Suffix"])

    result = compare_documents(left, right, include_reflow=False)
    text_anchors = [anchor for anchor in result.anchors if anchor.source_type == "text"]

    assert text_anchors
    assert all(anchor.confidence == "low" for anchor in text_anchors)


def test_compare_documents_merges_large_insert_region_into_page_bbox() -> None:
    inserted_words = [f"word{index}" for index in range(35)]
    left = _build_wrapped_document(["Prefix", "Suffix"])
    right = _build_wrapped_document(["Prefix", " ".join(inserted_words[:18]), " ".join(inserted_words[18:]), "Suffix"])

    result = compare_documents(left, right, include_reflow=False)
    insert_anchor = next(anchor for anchor in result.anchors if anchor.kind == "insert")

    assert insert_anchor.is_large_region
    assert len(insert_anchor.right_fragments) == 1
    assert insert_anchor.right_fragments[0].bbox is not None
    assert insert_anchor.right_fragments[0].bbox[3] > insert_anchor.right_fragments[0].bbox[1]


def test_compare_documents_emits_table_cell_anchor() -> None:
    left_path = Path(tempfile.gettempdir()) / "table-left.pdf"
    right_path = Path(tempfile.gettempdir()) / "table-right.pdf"

    _build_table_pdf(
        left_path,
        [
            ["Date", "Amount", "Status"],
            ["2025-04-15", "14,147.19", "成功"],
            ["2025-04-16", "9,000.00", "待处理"],
        ],
        top_label="outside-left",
    )
    _build_table_pdf(
        right_path,
        [
            ["Date", "Amount", "Status"],
            ["2025-04-15", "11,147.19", "成功"],
            ["2025-04-16", "9,000.00", "待处理"],
        ],
        top_label="outside-right",
    )

    from app.services.extractor import extract_document

    left = extract_document(left_path, header_margin=0, footer_margin=0)
    right = extract_document(right_path, header_margin=0, footer_margin=0)

    result = compare_documents(left, right, include_reflow=False)
    table_anchors = [anchor for anchor in result.anchors if anchor.source_type == "table"]

    assert len(table_anchors) == 1
    assert table_anchors[0].kind == "replace"
    assert table_anchors[0].confidence == "high"
    assert table_anchors[0].table_context is not None
    assert table_anchors[0].table_context.col_label == "Amount"
    assert "11,147.19" in table_anchors[0].excerpt_right


def test_compare_documents_does_not_mark_whole_table_replace_on_row_insert() -> None:
    left_path = Path(tempfile.gettempdir()) / "table-insert-left.pdf"
    right_path = Path(tempfile.gettempdir()) / "table-insert-right.pdf"

    _build_table_pdf(
        left_path,
        [
            ["Date", "Amount", "Status"],
            ["2025-04-15", "14,147.19", "成功"],
            ["2025-04-16", "9,000.00", "待处理"],
        ],
        top_label="outside-left",
    )
    _build_table_pdf(
        right_path,
        [
            ["Date", "Amount", "Status"],
            ["2025-04-15", "14,147.19", "成功"],
            ["2025-04-16", "9,000.00", "待处理"],
            ["2025-04-17", "1,000.00", "新增"],
        ],
        top_label="outside-right",
    )

    from app.services.extractor import extract_document

    left = extract_document(left_path, header_margin=0, footer_margin=0)
    right = extract_document(right_path, header_margin=0, footer_margin=0)

    result = compare_documents(left, right, include_reflow=False)
    structure_anchors = [
        anchor for anchor in result.anchors
        if anchor.source_type == "table" and anchor.table_context and anchor.table_context.row == -1
    ]
    inserted_cells = [
        anchor for anchor in result.anchors
        if anchor.source_type == "table" and anchor.kind == "insert"
    ]

    assert not structure_anchors
    assert inserted_cells
    assert all("2025-04-15" not in anchor.excerpt_right for anchor in inserted_cells)
