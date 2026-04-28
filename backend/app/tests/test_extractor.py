from pathlib import Path
import tempfile

from docx import Document as DocxDocument
import fitz

from app.services.extractor import SectionWindow, extract_document, slice_document_projection
from app.services.extractor import _clean_cell_text
from app.services.extractor import _should_keep_extracted_char
from app.services.normalizer import normalize_char


def test_normalize_char_preserves_length_for_spaces() -> None:
    assert normalize_char("\u00A0") == " "


def test_normalize_char_preserves_visible_character() -> None:
    assert normalize_char("A") == "A"


def test_extract_document_uses_real_char_boxes_from_rawdict() -> None:
    path = Path(tempfile.gettempdir()) / "extractor-rawdict-test.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 96), "Dose 0.05 mg")
    doc.save(path)
    doc.close()

    projection = extract_document(path, header_margin=0, footer_margin=0)

    visible_chars = [char for char in projection.chars if not char.synthetic]
    assert "".join(char.char for char in visible_chars) == "Dose 0.05 mg"
    assert all(char.bbox is not None for char in visible_chars)
    decimal_char = next(char for char in visible_chars if char.char == ".")
    assert decimal_char.bbox[2] > decimal_char.bbox[0]
    assert projection.words[1].text == "0.05"


def test_extract_document_splits_out_explicit_grid_tables() -> None:
    path = Path(tempfile.gettempdir()) / "extractor-table-test.pdf"
    doc = fitz.open()
    page = doc.new_page(width=320, height=220)
    page.insert_text((24, 24), "outside")

    x_positions = [24, 100, 180, 280]
    y_positions = [60, 96, 132]
    for x in x_positions:
        page.draw_line((x, y_positions[0]), (x, y_positions[-1]), color=(0, 0, 0), width=1)
    for y in y_positions:
        page.draw_line((x_positions[0], y), (x_positions[-1], y), color=(0, 0, 0), width=1)
    page.insert_text((32, 82), "A1")
    page.insert_text((108, 82), "B1")
    page.insert_text((188, 82), "C1")
    doc.save(path)
    doc.close()

    projection = extract_document(path, header_margin=0, footer_margin=0)

    assert len(projection.tables) == 1
    assert projection.tables[0].col_count == 3
    assert "outside" in projection.raw_text
    assert "A1" not in projection.raw_text


def test_extract_document_filters_repeated_running_headers() -> None:
    path = Path(tempfile.gettempdir()) / "extractor-running-header-test.pdf"
    doc = fitz.open()
    for page_index in range(3):
        page = doc.new_page(width=400, height=500)
        page.insert_text((36, 24), f"BeOne BGB-16673-304 A6 CONFIDENTIAL Page {page_index + 1}")
        page.insert_text((72, 160), f"Unique body line {page_index + 1}")
    doc.save(path)
    doc.close()

    projection = extract_document(path, header_margin=0, footer_margin=0)

    assert "CONFIDENTIAL" not in projection.raw_text
    assert "Unique body line 1" in projection.raw_text
    assert "Unique body line 2" in projection.raw_text
    assert "Unique body line 3" in projection.raw_text


def test_clean_cell_text_handles_none() -> None:
    assert _clean_cell_text(None) == ""
    assert _clean_cell_text("  Amount  ") == "Amount"


def test_should_keep_extracted_char_filters_control_chars() -> None:
    assert _should_keep_extracted_char("A")
    assert _should_keep_extracted_char(" ")
    assert not _should_keep_extracted_char("\x01")


def test_extract_document_builds_docx_projection_without_table_text_duplication() -> None:
    path = Path(tempfile.gettempdir()) / "extractor-docx-test.docx"
    document = DocxDocument()
    document.add_paragraph("Heading One", style="Heading 1")
    document.add_paragraph("Paragraph body text")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    document.save(path)

    projection = extract_document(path, header_margin=0, footer_margin=0)

    assert projection.document_kind == "docx"
    assert projection.pages == []
    assert "Heading One" in projection.raw_text
    assert "Paragraph body text" in projection.raw_text
    assert "A1" not in projection.raw_text
    assert len(projection.tables) == 1
    assert projection.tables[0].cells[0].dom_id == "table-2-r0-c0"
    assert projection.review_html is not None
    assert 'data-dom-id="heading-0"' in projection.review_html


def test_slice_document_projection_uses_same_page_y_window() -> None:
    path = Path(tempfile.gettempdir()) / "extractor-section-window-test.pdf"
    doc = fitz.open()
    page = doc.new_page(width=400, height=500)
    page.insert_text((72, 80), "Chapter heading")
    page.insert_text((72, 150), "section alpha")
    page.insert_text((72, 260), "section beta")
    doc.save(path)
    doc.close()

    projection = extract_document(path, header_margin=0, footer_margin=0)
    section = slice_document_projection(
        projection,
        SectionWindow(start_page=0, end_page=0, start_y=120, end_y=220),
    )

    assert "section alpha" in section.raw_text
    assert "Chapter heading" not in section.raw_text
    assert "section beta" not in section.raw_text
